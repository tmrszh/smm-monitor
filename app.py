"""Приложение «в один клик»: собирает отчёт и открывает его.

Запускается двойным щелчком (внутри — собранный PyInstaller исполняемый файл),
поэтому здесь нет аргументов командной строки: всё спрашивается у человека
и пишется на Рабочий стол.
"""

from __future__ import annotations

import json
import subprocess
import sys
import traceback
from datetime import date
from pathlib import Path

from smm_monitor.collectors import collect
from smm_monitor.net import CERTIFICATE_HELP, is_certificate_error, use_system_certificates
from smm_monitor.excel_report import build_workbook
from smm_monitor.attach import attach_counting
from smm_monitor.refill import refill_workbook
from smm_monitor.regions import RegionBook

DEFAULT_LIMIT = 100


def base_dir() -> Path:
    """Папка с данными: рядом с исполняемым файлом или рядом с исходниками."""
    if getattr(sys, "frozen", False):
        return Path(sys._MEIPASS)  # type: ignore[attr-defined]
    return Path(__file__).resolve().parent


def windows_desktop() -> Path | None:
    """Настоящий Рабочий стол Windows: при включённом OneDrive он лежит внутри него."""
    try:
        import ctypes
        from ctypes import windll, wintypes

        class GUID(ctypes.Structure):
            _fields_ = [
                ("Data1", wintypes.DWORD),
                ("Data2", wintypes.WORD),
                ("Data3", wintypes.WORD),
                ("Data4", ctypes.c_byte * 8),
            ]

        desktop_id = GUID(
            0xB4BFCC3A, 0xDB2C, 0x424C,
            (ctypes.c_byte * 8)(0xB0, 0x29, 0x7F, 0xE9, 0x9A, 0x87, 0xC6, 0x41),
        )
        buffer = ctypes.c_wchar_p()
        if windll.shell32.SHGetKnownFolderPath(
            ctypes.byref(desktop_id), 0, None, ctypes.byref(buffer)
        ) != 0:
            return None
        path = Path(buffer.value)
        windll.ole32.CoTaskMemFree(buffer)
        return path if path.is_dir() else None
    except Exception:
        return None


def app_folder() -> Path:
    """Папка, где лежит сам файл программы (туда же можно положить сертификат)."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def use_local_certificate() -> Path | None:
    """Если рядом с программой лежит файл сертификата организации — взять его."""
    import os

    for name in ("сертификат.pem", "certificate.pem", "сертификат.crt", "certificate.crt"):
        path = app_folder() / name
        if path.is_file():
            os.environ["REQUESTS_CA_BUNDLE"] = str(path)
            os.environ["SSL_CERT_FILE"] = str(path)
            return path
    return None


MEMORY_FILE = "последний-файл.txt"


def remember_file(path: Path) -> None:
    try:
        (app_folder() / MEMORY_FILE).write_text(str(path), encoding="utf-8")
    except Exception:
        pass


def remembered_file() -> Path | None:
    try:
        path = Path((app_folder() / MEMORY_FILE).read_text(encoding="utf-8").strip())
        return path if path.is_file() else None
    except Exception:
        return None


def nearby_workbook() -> Path | None:
    """Таблица с листом «Ссылки», лежащая рядом с программой или на Рабочем столе."""
    from openpyxl import load_workbook

    folders = [app_folder(), desktop()]
    candidates = []
    for folder in folders:
        try:
            candidates.extend(sorted(folder.glob("*.xlsx"), key=lambda f: -f.stat().st_mtime))
        except Exception:
            continue
    standard = {"Ссылки", "Регионы", "Справочник"}
    best: tuple[int, Path] | None = None
    for candidate in candidates[:20]:
        if candidate.name.startswith("~$"):
            continue
        try:
            workbook = load_workbook(candidate, read_only=True)
            sheets = set(workbook.sheetnames)
            workbook.close()
        except Exception:
            continue
        if "Ссылки" not in sheets:
            continue
        # таблица самого человека (есть свой лист) важнее пустой заготовки
        score = 2 if sheets - standard else 1
        if best is None or score > best[0]:
            best = (score, candidate)
    return best[1] if best else None


def desktop() -> Path:
    if sys.platform.startswith("win"):
        found = windows_desktop()
        if found:
            return found
        one_drive = Path.home() / "OneDrive" / "Desktop"
        if one_drive.is_dir():
            return one_drive
    for name in ("Desktop", "Рабочий стол"):
        path = Path.home() / name
        if path.is_dir():
            return path
    return Path.home()


def open_file(path: Path) -> None:
    try:
        if sys.platform == "darwin":
            subprocess.run(["open", str(path)], check=False)
        elif sys.platform.startswith("win"):
            import os

            os.startfile(str(path))  # type: ignore[attr-defined]
        else:
            subprocess.run(["xdg-open", str(path)], check=False)
    except Exception:
        pass


def ask_limit() -> int:
    print("Сколько последних публикаций брать с каждой площадки?")
    answer = input(f"Нажмите Enter для {DEFAULT_LIMIT} или введите число: ").strip()
    if not answer:
        return DEFAULT_LIMIT
    try:
        value = int(answer)
    except ValueError:
        print(f"Не похоже на число, беру {DEFAULT_LIMIT}.")
        return DEFAULT_LIMIT
    return max(1, min(value, 1000))


def ask_mode(known: Path | None) -> str:
    print("Что сделать?")
    if known:
        print(f"  Enter — обновить цифры в таблице «{known.name}»")
        print("  1 — собрать свежий отчёт по каналам АФМ")
        print("  3 — настроить автоподсчёт в другой таблице")
        hint = "Нажмите Enter или введите 1 или 3: "
        default = "2"
    else:
        print("  1 — собрать свежий отчёт по каналам АФМ")
        print("  2 — дополнить файл, в который вы вставили ссылки")
        print("  3 — настроить автоподсчёт в вашей таблице (делается один раз)")
        hint = "Введите 1, 2 или 3 и нажмите Enter: "
        default = "1"
    print()
    while True:
        answer = input(hint).strip()
        if answer in ("1", "2", "3", ""):
            return answer or default
        print("Нужно 1, 2 или 3.")


def ask_file(suggest: Path | None = None) -> Path | None:
    print()
    if suggest:
        print(f"Файл: {suggest.name}")
        raw = input("Нажмите Enter, чтобы взять его, или перетащите сюда другой: ").strip()
        if not raw:
            return suggest
    else:
        print("Перетащите файл Excel в это окно и нажмите Enter.")
        raw = input("Путь к файлу: ").strip()
    raw = raw.strip().strip("'\"")
    if not raw:
        return None
    path = Path(raw.replace("\\ ", " ")).expanduser()
    if not path.is_file():
        print(f"Файл не найден: {path}")
        return None
    return path


def own_pages() -> list[str]:
    try:
        config = json.loads((base_dir() / "config" / "sources.json").read_text(encoding="utf-8"))
        return config.get("own_pages") or []
    except Exception:
        return []


def run_refill(region_book: RegionBook, known: Path | None = None) -> None:
    path = ask_file(known or remembered_file() or nearby_workbook())
    if path is None:
        return
    remember_file(path)
    print()
    print("Иду по ссылкам, это может занять несколько минут. Окно не закрывайте.")
    print()
    result = refill_workbook(
        path,
        region_book,
        progress=lambda link: print(f"  {link}", flush=True),
        own_pages=own_pages(),
    )
    if any(is_certificate_error(Exception(problem)) for problem in result.failed):
        print(CERTIFICATE_HELP)
    print()
    print(f"Заполнено строк: {result.filled}")
    if result.foreign:
        print(
            f"Из них ссылок на другие ресурсы: {result.foreign} — "
            "они посчитаны как публикации, без просмотров и комментариев"
        )
    if result.skipped:
        print(f"Пропущено, цифры уже были: {result.skipped}")
    for problem in result.failed:
        print(f"Не получилось: {problem}")
    print()
    print(f"Файл сохранён: {path}")
    open_file(path)


NAMES = {"links": "репосты (количество публикаций)", "views": "просмотры", "comments": "комментарии", "place": "место"}


def run_attach(region_book: RegionBook) -> None:
    print()
    print("Нужна ваша таблица с регионами. Программа не меняет её, а делает копию:")
    print("добавляет лист «Ссылки» и ставит формулы в колонки «кол-во».")
    path = ask_file()
    if path is None:
        return
    try:
        out, result = attach_counting(path, region_book)
        remember_file(out)
    except Exception as error:
        print()
        print(f"Не получилось: {error}")
        return

    print()
    if result.already_prepared:
        print("Эта таблица уже настроена, ничего менять не нужно.")
        print(f"Файл: {out}")
        print()
        print("Чтобы обновить цифры, вставьте ссылки на лист «Ссылки»,")
        print("сохраните файл и запустите программу ещё раз — пункт 2.")
        open_file(out)
        return
    print(f"Готово: {out}")
    print(f"Лист «{result.sheet}», регионов: {len(result.regions)}")
    print("Считаются сами: " + ", ".join(NAMES[m] for m in result.columns))
    if result.missing:
        print("Не нашлось колонок: " + ", ".join(NAMES[m] for m in result.missing))
    print()
    print("Дальше: вставьте ссылки на лист «Ссылки» в колонку «Ссылка»,")
    print("сохраните файл, запустите программу и выберите пункт 2.")
    open_file(out)


def run_collect(sources: list, region_book: RegionBook) -> None:
    limit = ask_limit()
    print()
    print("Собираю. Телеграм занимает около минуты, инстаграм — несколько минут.")
    print("Окно не закрывайте.")
    print()

    posts = []
    for source in sources:
        title = source.get("title") or source["url"]
        print(f"  {title} ...", flush=True)
        try:
            found = collect(source, limit, region_book)
        except Exception as error:
            if is_certificate_error(error):
                print(f"    {CERTIFICATE_HELP}")
            else:
                print(f"    не удалось собрать: {error}")
            continue
        print(f"    публикаций: {len(found)}")
        posts.extend(found)

    if not posts:
        print()
        print("Ничего не собралось. Проверьте интернет и попробуйте ещё раз.")
        return

    out = desktop() / f"Отчёт АФМ {date.today():%d.%m.%Y}.xlsx"
    build_workbook(posts, region_book, out)

    print()
    print(f"Готово: {out}")
    print("Файл сохранён на Рабочий стол и сейчас откроется.")
    open_file(out)


def main() -> None:
    print("=" * 58)
    print("  Отчёт по соцсетям АФМ")
    print("=" * 58)
    print()

    use_system_certificates()
    certificate = use_local_certificate()
    if certificate:
        print(f"Использую сертификат организации: {certificate.name}")
        print()

    config = json.loads((base_dir() / "config" / "sources.json").read_text(encoding="utf-8"))
    region_book = RegionBook.load(base_dir() / "config" / "regions.json")

    known = remembered_file() or nearby_workbook()
    mode = ask_mode(known)
    if mode == "2":
        run_refill(region_book, known)
    elif mode == "3":
        run_attach(region_book)
    else:
        run_collect(config["sources"], region_book)


if __name__ == "__main__":
    try:
        main()
    except Exception:
        print()
        print("Что-то пошло не так:")
        traceback.print_exc()
    finally:
        print()
        input("Нажмите Enter, чтобы закрыть окно...")
