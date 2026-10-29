import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from smm_monitor.excel_report import build_workbook
from smm_monitor.regions import RegionBook
from tests.demo_posts import demo_posts


def main() -> None:
    book = RegionBook.load()
    posts = demo_posts(book)
    assert posts[0].regions == ["Туркестан"], posts[0].regions
    assert posts[1].regions == ["Туркестан", "г. Шымкент"], posts[1].regions
    assert posts[5].regions == []

    path = build_workbook(posts, book, Path("out/demo-report.xlsx"))

    # пересчёт формул настоящим Excel не делаем — проверяем, что они на месте
    from openpyxl import load_workbook

    wb = load_workbook(path)
    regions = wb["Регионы"]
    row = next(r for r in range(2, 40) if regions.cell(row=r, column=1).value == "Туркестан")
    formula = regions.cell(row=row, column=2).value
    assert formula.startswith("=COUNTIF('Ссылки'!$E$2:$E$1001"), formula
    assert wb["Ссылки"].cell(row=2, column=5).value == "Туркестан"
    print("OK:", path.resolve())
    print("Туркестан, формула количества:", formula)


if __name__ == "__main__":
    main()
