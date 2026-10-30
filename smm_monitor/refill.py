"""Дозаполнение готового файла: по вставленным вручную ссылкам подтянуть метрики.

Человек вставляет ссылки в колонку «Ссылка» на листе «Ссылки» и запускает сбор.
Пустые ячейки строки заполняются, уже заполненные руками не трогаются
(кроме режима обновления, когда цифры перезаписываются свежими).
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

from .collectors.single import fetch_post
from .excel_report import LINKS_SHEET
from .regions import RegionBook

COL_NUMBER, COL_PLATFORM, COL_LINK, COL_DATE = 1, 2, 3, 4
COL_REGION, COL_DESCRIPTION, COL_VIEWS, COL_COMMENTS = 5, 6, 7, 8


@dataclass
class RefillResult:
    filled: int = 0
    updated: int = 0
    skipped: int = 0
    failed: list[str] = None  # type: ignore[assignment]

    def __post_init__(self) -> None:
        if self.failed is None:
            self.failed = []


def refill_workbook(
    path: str | Path,
    region_book: RegionBook,
    refresh: bool = False,
    progress=None,
) -> RefillResult:
    path = Path(path)
    workbook = load_workbook(path)
    if LINKS_SHEET not in workbook.sheetnames:
        raise ValueError(f"В файле нет листа «{LINKS_SHEET}»")
    sheet = workbook[LINKS_SHEET]
    result = RefillResult()

    for row in range(2, sheet.max_row + 1):
        link = sheet.cell(row=row, column=COL_LINK).value
        if not link or not str(link).strip().startswith("http"):
            continue
        link = str(link).strip()

        has_numbers = any(
            sheet.cell(row=row, column=column).value is not None
            for column in (COL_VIEWS, COL_COMMENTS)
        )
        if has_numbers and not refresh:
            result.skipped += 1
            continue

        if progress:
            progress(link)
        try:
            post = fetch_post(link)
        except Exception as error:
            result.failed.append(f"{link} — {error}")
            continue
        if post is None:
            result.failed.append(f"{link} — ссылка не распознана")
            continue

        def put(column: int, value, overwrite: bool = False) -> None:
            cell = sheet.cell(row=row, column=column)
            if value is not None and (overwrite or cell.value in (None, "")):
                cell.value = value

        put(COL_NUMBER, row - 1)
        put(COL_PLATFORM, post.platform.capitalize())
        put(COL_DATE, (post.posted_at or "")[:19].replace("T", " ") or None)
        put(COL_DESCRIPTION, post.description or None)
        put(COL_REGION, "; ".join(region_book.detect(post.description)) or None)
        for column, value in ((COL_VIEWS, post.views), (COL_COMMENTS, post.comments)):
            put(column, value, overwrite=refresh)

        if has_numbers:
            result.updated += 1
        else:
            result.filled += 1

    workbook.save(path)
    return result
