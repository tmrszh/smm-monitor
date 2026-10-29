"""Сборка Excel-отчёта.

Лист «Ссылки»   — по строке на каждый пост: ссылка, описание, регион и метрики.
Лист «Регионы»  — таблица регионов, где напротив каждого региона живыми формулами
                  считается количество ссылок и сумма просмотров/лайков/комментариев.
Лист «Справочник» — список регионов (источник выпадающего списка в «Ссылках»).

Формулы ссылаются на диапазон до MAX_ROWS строк, поэтому добавленные руками ссылки
пересчитываются сами — пересобирать файл не нужно.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Post
from .regions import RegionBook

LINKS_SHEET = "Ссылки"
REGIONS_SHEET = "Регионы"
BOOK_SHEET = "Справочник"

MAX_ROWS = 1000  # запас строк, на которые рассчитаны формулы

LINK_COLUMNS = [
    ("№", 6),
    ("Платформа", 13),
    ("Ссылка", 46),
    ("Дата", 18),
    ("Регион", 22),
    ("Описание", 60),
    ("Просмотры", 12),
    ("Лайки", 10),
    ("Комментарии", 13),
]
COL_PLATFORM = "B"
COL_LINK = "C"
COL_REGION = "E"
COL_VIEWS = "G"
COL_LIKES = "H"
COL_COMMENTS = "I"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)


def _rng(column: str) -> str:
    return f"'{LINKS_SHEET}'!${column}$2:${column}${MAX_ROWS + 1}"


def _style_header(sheet, columns) -> None:
    for index, (title, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}1"


def _fill_links(sheet, posts: list[Post], region_book: RegionBook) -> None:
    _style_header(sheet, LINK_COLUMNS)

    for offset, post in enumerate(posts):
        row = offset + 2
        sheet.cell(row=row, column=1, value=offset + 1)
        sheet.cell(row=row, column=2, value=post.platform.capitalize())
        link = sheet.cell(row=row, column=3, value=post.url)
        link.hyperlink = post.url
        link.style = "Hyperlink"
        sheet.cell(row=row, column=4, value=(post.posted_at or "")[:19].replace("T", " "))
        sheet.cell(row=row, column=5, value=post.regions_cell)
        description = sheet.cell(row=row, column=6, value=post.description)
        description.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(row=row, column=7, value=post.views)
        sheet.cell(row=row, column=8, value=post.likes)
        sheet.cell(row=row, column=9, value=post.comments)

    # выпадающий список регионов для колонки «Регион» (на весь запас строк)
    formula = f"='{BOOK_SHEET}'!$A$2:$A${len(region_book.names) + 1}"
    validation = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    validation.error = "Выберите регион из справочника"
    validation.errorTitle = "Регион не из списка"
    sheet.add_data_validation(validation)
    validation.add(f"{COL_REGION}2:{COL_REGION}{MAX_ROWS + 1}")


def _fill_regions(sheet, region_book: RegionBook) -> None:
    columns = [
        ("Регион", 24),
        ("Ссылок", 11),
        ("Instagram", 11),
        ("Telegram", 11),
        ("Просмотры", 13),
        ("Лайки", 11),
        ("Комментарии", 13),
    ]
    _style_header(sheet, columns)

    region_rng, platform_rng, link_rng = _rng(COL_REGION), _rng(COL_PLATFORM), _rng(COL_LINK)

    for offset, name in enumerate(region_book.names):
        row = offset + 2
        key = f'"*"&$A{row}&"*"'
        sheet.cell(row=row, column=1, value=name)
        sheet.cell(row=row, column=2, value=f"=COUNTIF({region_rng},{key})")
        sheet.cell(row=row, column=3, value=f'=COUNTIFS({platform_rng},"Instagram",{region_rng},{key})')
        sheet.cell(row=row, column=4, value=f'=COUNTIFS({platform_rng},"Telegram",{region_rng},{key})')
        sheet.cell(row=row, column=5, value=f"=SUMIF({region_rng},{key},{_rng(COL_VIEWS)})")
        sheet.cell(row=row, column=6, value=f"=SUMIF({region_rng},{key},{_rng(COL_LIKES)})")
        sheet.cell(row=row, column=7, value=f"=SUMIF({region_rng},{key},{_rng(COL_COMMENTS)})")

    last = len(region_book.names) + 1

    no_region = last + 1
    sheet.cell(row=no_region, column=1, value="Без региона").font = TOTAL_FONT
    sheet.cell(
        row=no_region,
        column=2,
        value=f'=COUNTIFS({link_rng},"<>",{region_rng},"")',
    ).font = TOTAL_FONT

    total = last + 2
    sheet.cell(row=total, column=1, value="Всего ссылок").font = TOTAL_FONT
    sheet.cell(row=total, column=2, value=f"=COUNTA({link_rng})").font = TOTAL_FONT
    for column, source in ((5, COL_VIEWS), (6, COL_LIKES), (7, COL_COMMENTS)):
        sheet.cell(row=total, column=column, value=f"=SUM({_rng(source)})").font = TOTAL_FONT

    note = total + 2
    sheet.cell(
        row=note,
        column=1,
        value=(
            f"Формулы считают диапазон «{LINKS_SHEET}» до {MAX_ROWS} строк: "
            "новые ссылки достаточно вставить на лист «Ссылки», цифры пересчитаются сами. "
            "Пост с несколькими регионами попадает в каждый из них."
        ),
    ).alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(start_row=note, start_column=1, end_row=note + 2, end_column=7)


def _fill_book(sheet, region_book: RegionBook) -> None:
    _style_header(sheet, [("Регион", 24), ("Ключевые слова для авто-определения", 70)])
    for offset, region in enumerate(region_book.regions):
        row = offset + 2
        sheet.cell(row=row, column=1, value=region.name)
        sheet.cell(row=row, column=2, value=", ".join(region.aliases))


def build_workbook(posts: list[Post], region_book: RegionBook, path: str | Path) -> Path:
    workbook = Workbook()
    links = workbook.active
    links.title = LINKS_SHEET
    regions = workbook.create_sheet(REGIONS_SHEET)
    book = workbook.create_sheet(BOOK_SHEET)

    _fill_links(links, posts, region_book)
    _fill_regions(regions, region_book)
    _fill_book(book, region_book)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path
